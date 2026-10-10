import json
import time
import openpyxl
from duckduckgo_search import DDGS

# Load matched items
with open('matched_768_items.json', 'r', encoding='utf-8') as f:
    items = json.load(f)

file_path = 'HRZ_Pitstop_Client_Inventory_All_Bike_Models.xlsx'
wb = openpyxl.load_workbook(file_path)

print("Starting verification against Legundary and Bandidos Pitstop...")
ddgs = DDGS()

def get_online_verification(name):
    query = f"{name} site:legundary.com OR site:bandidospitstop.com"
    try:
        results = list(ddgs.text(query, max_results=2))
        if not results: return None, None
        url = str(results[0].get('href', '')).lower()
        title = str(results[0].get('title', ''))
        if 'legundary.com' in url: return 'Legundary Customs', title
        if 'bandidospitstop.com' in url: return 'Bandidos Pitstop', title
    except Exception as e:
        return None, None
    return None, None

# Load Site Products to get the site names mapped
with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)
    
site_map = {p['name'].lower(): p for p in products}

verified_count = 0

# Iterate through sheets and update Brand if verified
for sheet_name in wb.sheetnames[:14]:
    sheet = wb[sheet_name]
    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row, column=1).value
        current_brand = sheet.cell(row=row, column=3).value
        
        if not name: continue
        
        # If it's already Legundary from local data, we can consider it verified.
        if current_brand and 'legundary' in str(current_brand).lower():
            verified_count += 1
            sheet.cell(row=row, column=3, value='Legundary Customs (Verified)')
            continue
            
        # Try to search DDGS
        print(f"Searching online for: {name}")
        brand_found, verified_title = get_online_verification(name)
        if brand_found:
            sheet.cell(row=row, column=3, value=f"{brand_found} (Verified Online)")
            verified_count += 1
            print(f"-> Found! {brand_found}")
        else:
            print("-> Not found on those sites.")
            sheet.cell(row=row, column=3, value=f"{current_brand} (Unverified)")
            
        time.sleep(1) # Be nice to the API to avoid rate limits

wb.save(file_path)
print(f"Verification complete. Total items verified from Legundary/Bandidos: {verified_count}")
