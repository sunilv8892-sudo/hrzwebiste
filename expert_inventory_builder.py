import json
import openpyxl
import re

def get_expert_category(name):
    n = str(name).lower()
    
    # 1. Helmets & Accessories
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'helmet', 'helmets', 'visor', 'pinlock', 'goggle', 'goggles', 'lens', 'intercom', 'bluetooth', 
        'helmet lock', 'helmet mount', 'spoiler', 'anti-fog', 'anti fog', 'breath deflector', 'helmet cap', 'helmet bag'
    ]):
        return 'Helmets & Accessories'
        
    # 2. Riding Gears
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'jacket', 'pant', 'pants', 'glove', 'gloves', 'suit', 'shoe', 'shoes', 'boot', 'boots', 'jersey', 
        'knee guard', 'elbow guard', 'rain', 'rainwear', 'raincover', 'bionic', 'armor', 'armour', 'base layer',
        'balaclava', 'socks', 'riding denim', 'riding jeans', 'protector', 'spine guard', 'riding jacket'
    ]):
        return 'Riding Gears'
        
    # 3. Luggage & Touring
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'bag', 'bags', 'luggage', 'topcase', 'top case', 'saddle', 'saddlebag', 'saddlebags', 'pannier', 'panniers', 
        'carrier', 'tail bag', 'tank bag', 'backpack', 'backrest', 'back rest', 'hydration', 'bladder', 
        'touring rack', 'rear rack', 'bungee', 'cargo net'
    ]):
        return 'Luggage & Touring'
        
    # 4. Lighting
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'light', 'lights', 'fog', 'indicator', 'indicators', 'led', 'bulb', 'flasher', 'hazard', 'headlight', 
        'head lamp', 'tail light', 'auxiliary', 'switch', 'wire harness', 'wiring', 'relay'
    ]):
        return 'Lighting'
        
    # 5. Bike Protection
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'crashguard', 'crash guard', 'slider', 'sliders', 'bash plate', 'bashplate', 'sump guard', 'engine guard', 
        'radiator guard', 'radiator grill', 'leg guard', 'frame slider', 'tank pad', 'fork protector', 
        'exhaust protector', 'caliper guard'
    ]):
        return 'Bike Protection'
        
    # 6. Handlebar & Accessories
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'handle', 'handlebar', 'grip', 'grips', 'lever', 'levers', 'mirror', 'mirrors', 'bar end', 'bar ends', 
        'riser', 'risers', 'knuckle', 'knuckle guard', 'handguard', 'hand guard', 'throttle'
    ]):
        return 'Handlebar & Accessories'
        
    # 7. Foot Rests
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'footpeg', 'foot peg', 'footpegs', 'footrest', 'foot rest', 'foot rests', 'heel guard', 'toe guard'
    ]):
        return 'Foot Rests'
        
    # 8. Guards & Mounts (For non-protection mounts, GPS, master cylinder guards, etc)
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'mount', 'bracket', 'holder', 'stay', 'gps', 'mobile holder', 'phone mount', 'camera mount', 
        'master cylinder guard', 'fluid reservoir guard', 'clamp'
    ]):
        return 'Guards & Mounts'
        
    # 9. Oil Caps & Stands
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'stand', 'stands', 'paddock', 'oil cap', 'spool', 'spools', 'side stand base', 'center stand', 'main stand',
        'fluid cap', 'reservoir cap'
    ]):
        return 'Oil Caps & Stands'
        
    # 10. Body Fairing and Fenders
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'fairing', 'fender', 'mudguard', 'winglet', 'winglets', 'cowl', 'windscreen', 'windshield', 'tail tidy',
        'number plate', 'license plate', 'visor \(bike\)', 'aerodynamic'
    ]):
        return 'Body Fairing and Fenders'
        
    # 11. Cleaning Accessories
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'cleaner', 'brush', 'lube', 'polish', 'cloth', 'micro fiber', 'microfiber', 'wash', 'chain lube',
        'chain clean', 'degreaser', 'wax', 'cleaning kit'
    ]):
        return 'Cleaning Accessories'
        
    # 12. Performance Accessories
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'filter', 'air filter', 'oil filter', 'exhaust', 'spark plug', 'chain sprocket', 'sprocket', 'chain kit',
        'brake pad', 'brake pads', 'clutch plate', 'suspension', 'performance'
    ]):
        return 'Performance Accessories'
        
    # 13. Bike Essentials
    if any(re.search(r'\b{}\b'.format(x), n) for x in [
        'cover', 'bike cover', 'horn', 'cable', 'charger', 'usb', 'tyre', 'tire', 'tube', 'battery', 
        'engine oil', 'coolant', 'brake fluid'
    ]):
        return 'Bike Essentials'
        
    # 14. Default
    return 'General Accessories'

def clean_text(text):
    text = str(text).lower()
    text = re.sub(r'[^a-z0-9\s]', ' ', text)
    return ' '.join(text.split())

# 1. Load exact online product details
with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)

# Build map by exact name and "brand + name"
site_products_map = {}
for p in products:
    b = p.get('brand', '')
    n = p.get('name', '')
    site_products_map[clean_text(f"{b} {n}")] = p
    site_products_map[clean_text(n)] = p

# 2. Load the Excel file and wipe existing data safely
file_path = 'HRZ_Pitstop_Client_Inventory_All_Bike_Models.xlsx'
wb = openpyxl.load_workbook(file_path)
main_sheets = wb.sheetnames[:14]

for s_name in main_sheets:
    sheet = wb[s_name]
    # Keep row 1 (headers), delete values from row 2 down
    for row in range(2, sheet.max_row + 1):
        for col in range(1, 13):
            sheet.cell(row=row, column=col).value = None

# 3. Process matched items and add to excel
with open('matched_768_items.json', 'r', encoding='utf-8') as f:
    items = json.load(f)

row_pointers = {s: 2 for s in main_sheets}
counts = {s: 0 for s in main_sheets}

for item in items:
    raw_name = item['raw_name']
    cleaned = clean_text(raw_name)
    
    # Resolve the site product strictly
    site_product = site_products_map.get(cleaned)
    if not site_product:
        # Fallback to high-confidence subset match
        tokens = set(cleaned.split())
        for sp_name, sp_obj in site_products_map.items():
            st = set(sp_name.split())
            if len(tokens) > 1 and len(tokens & st) / len(tokens) >= 0.8:
                site_product = sp_obj
                break
                
    if not site_product:
        continue
        
    # Get the official online name and categorize
    site_name = site_product.get('name', raw_name)
    category = get_expert_category(site_name)
    
    if category not in main_sheets:
        category = 'General Accessories'
        
    sheet = wb[category]
    curr_row = row_pointers[category]
    
    # Gather exact data from the JSON database (online verified values)
    brand = site_product.get('brand', '')
    mrp = site_product.get('originalPrice', '')
    sp = site_product.get('price', '')
    stock = item['excel_row'].get('Available Quantity for Sale', 10) # Prefer raw stock if available
    
    compats = site_product.get('bike_compatibility', [])
    bike_make = ''
    bike_model = ''
    if compats and len(compats) > 0:
        parts = compats[0].split(' ')
        bike_make = parts[0]
        if len(parts) > 1:
            bike_model = ' '.join(parts[1:])
    
    image = site_product.get('image', '')
    desc = site_product.get('description', '')
    if len(desc) > 150: desc = desc[:147] + '...'
    
    discount = ''
    if mrp and sp and mrp > 0:
        discount = round(((mrp - sp) / mrp) * 100)
    
    # 4. Populate row
    sheet.cell(row=curr_row, column=1, value=site_name)
    sheet.cell(row=curr_row, column=2, value=category) # Subcategory matches target category sheet
    sheet.cell(row=curr_row, column=3, value=brand)
    sheet.cell(row=curr_row, column=4, value=bike_make)
    sheet.cell(row=curr_row, column=5, value=bike_model)
    sheet.cell(row=curr_row, column=6, value=mrp)
    sheet.cell(row=curr_row, column=7, value=sp)
    sheet.cell(row=curr_row, column=8, value=discount)
    sheet.cell(row=curr_row, column=9, value=stock)
    sheet.cell(row=curr_row, column=10, value=image)
    sheet.cell(row=curr_row, column=11, value=desc)

    row_pointers[category] += 1
    counts[category] += 1

wb.save(file_path)
print("Inventory populated successfully using expert knowledge logic.")
print("Category breakdown:")
for k, v in counts.items():
    if v > 0:
        print(f"  {k}: {v} items")
