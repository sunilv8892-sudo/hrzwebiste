import openpyxl
import re

file_path = 'HRZ_Pitstop_Client_Inventory_All_Bike_Models.xlsx'
wb = openpyxl.load_workbook(file_path)

def determine_subcategory(sheet_name, product_name):
    n = str(product_name).lower()
    
    if sheet_name == 'Bike Protection':
        if 'crash' in n: return 'Crash Guard'
        if 'bash' in n or 'sump' in n: return 'Bash Plate'
        if 'slider' in n: return 'Frame Slider'
        if 'radiator' in n: return 'Radiator Guard'
        if 'caliper' in n: return 'Caliper Guard'
        if 'leg guard' in n: return 'Leg Guard'
        return 'Protection Accessories'
        
    if sheet_name == 'Luggage & Touring':
        if 'top case' in n or 'topcase' in n: return 'Top Case'
        if 'saddle' in n: return 'Saddlebag'
        if 'pannier' in n: return 'Pannier'
        if 'carrier' in n or 'rack' in n: return 'Carrier / Rack'
        if 'backrest' in n or 'back rest' in n: return 'Backrest'
        if 'tank' in n: return 'Tank Bag'
        if 'tail' in n: return 'Tail Bag'
        return 'Touring Accessories'
        
    if sheet_name == 'Helmets & Accessories':
        if 'visor' in n or 'shield' in n: return 'Visor'
        if 'pinlock' in n: return 'Pinlock'
        if 'intercom' in n or 'bluetooth' in n: return 'Intercom'
        if 'helmet' in n and ('bag' not in n and 'lock' not in n): return 'Helmet'
        return 'Helmet Accessories'
        
    if sheet_name == 'Riding Gears':
        if 'jacket' in n: return 'Riding Jacket'
        if 'pant' in n: return 'Riding Pants'
        if 'glove' in n: return 'Riding Gloves'
        if 'boot' in n or 'shoe' in n: return 'Riding Boots'
        if 'rain' in n: return 'Rainwear'
        if 'knee' in n or 'elbow' in n or 'bionic' in n: return 'Armour & Guards'
        return 'Riding Apparel'
        
    if sheet_name == 'Lighting':
        if 'fog' in n: return 'Fog Lights'
        if 'indicator' in n: return 'Indicators'
        if 'headlight' in n or 'bulb' in n: return 'Headlights / Bulbs'
        if 'switch' in n or 'harness' in n or 'relay' in n: return 'Wiring & Switches'
        return 'Lighting Accessories'
        
    if sheet_name == 'Handlebar & Accessories':
        if 'grip' in n: return 'Handle Grips'
        if 'lever' in n: return 'Levers'
        if 'mirror' in n: return 'Mirrors'
        if 'riser' in n: return 'Handlebar Risers'
        if 'knuckle' in n or 'hand guard' in n: return 'Handguards'
        return 'Handlebar Accessories'
        
    if sheet_name == 'Foot Rests':
        if 'peg' in n or 'rest' in n: return 'Footpegs'
        return 'Foot Rests'
        
    if sheet_name == 'Guards & Mounts':
        if 'gps' in n or 'mobile' in n or 'phone' in n: return 'Mobile / GPS Mount'
        if 'stay' in n: return 'Stays & Brackets'
        return 'Mounts'
        
    if sheet_name == 'Oil Caps & Stands':
        if 'paddock' in n: return 'Paddock Stand'
        if 'side stand' in n: return 'Side Stand Base'
        if 'cap' in n: return 'Fluid / Oil Caps'
        if 'spool' in n: return 'Swingarm Spools'
        return 'Stands & Caps'
        
    if sheet_name == 'Body Fairing and Fenders':
        if 'winglet' in n: return 'Winglets'
        if 'windscreen' in n or 'visor' in n: return 'Windscreen'
        if 'tail tidy' in n or 'number plate' in n: return 'Tail Tidy'
        if 'mudguard' in n or 'fender' in n: return 'Fender'
        return 'Body Kits'
        
    if sheet_name == 'Performance Accessories':
        if 'filter' in n: return 'Air / Oil Filters'
        if 'exhaust' in n: return 'Exhausts'
        if 'sprocket' in n or 'chain kit' in n: return 'Chain & Sprockets'
        if 'spark plug' in n: return 'Spark Plugs'
        return 'Performance Parts'
        
    if sheet_name == 'Cleaning Accessories':
        if 'chain' in n and ('lube' in n or 'clean' in n): return 'Chain Lube/Cleaner'
        if 'polish' in n or 'wax' in n: return 'Polish & Wax'
        if 'brush' in n or 'cloth' in n: return 'Cleaning Tools'
        return 'Bike Care'
        
    if sheet_name == 'Bike Essentials':
        if 'cover' in n: return 'Bike Cover'
        if 'horn' in n: return 'Horns'
        if 'charger' in n: return 'Chargers'
        return 'Essentials'

    return 'General'

main_sheets = wb.sheetnames[:14]
updated_count = 0

for s_name in main_sheets:
    sheet = wb[s_name]
    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row, column=1).value
        if not name: continue
        
        # Calculate proper subcategory based on name
        new_subcat = determine_subcategory(s_name, name)
        
        # Update Column 2
        sheet.cell(row=row, column=2, value=new_subcat)
        updated_count += 1

wb.save(file_path)
print(f"Subcategories fixed. Updated {updated_count} rows with granular subcategories.")
