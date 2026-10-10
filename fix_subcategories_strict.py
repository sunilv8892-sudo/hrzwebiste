import openpyxl

file_path = 'HRZ_Pitstop_Client_Inventory_All_Bike_Models.xlsx'
wb = openpyxl.load_workbook(file_path)

def determine_subcategory_strict(sheet_name, product_name):
    n = str(product_name).lower()
    
    if sheet_name == 'Bike Protection':
        if 'crash' in n or 'leg guard' in n: return 'Crash Guard'
        if 'slider' in n: return 'Frame Slider'
        if 'master cylinder' in n: return 'Master Cylinder Guard'
        if 'bash' in n or 'sump' in n or 'engine guard' in n or 'skid plate' in n: return 'Engine Guards and Skid Plate'
        if 'headlight' in n and 'grill' in n: return 'Headlight Grill'
        if 'radiator' in n: return 'Radiator Guard'
        if 'tank' in n and 'pad' in n: return 'Tank Protectors'
        if 'stand' in n: return 'Stand and Extenders'
        if 'chain' in n or 'sprocket' in n: return 'Chain Cover And Sprocket'
        if 'cap' in n or 'fluid' in n: return 'Fluid Guards and Caps'
        return 'Crash Guard' # Default fallback
        
    if sheet_name == 'Handlebar & Accessories':
        if 'grip' in n or 'throttle' in n: return 'Grips and Throttle'
        if 'hand guard' in n or 'knuckle' in n: return 'Hand Guard'
        if 'riser' in n: return 'Handle Risers'
        if 'lever' in n: return 'Lever Guard'
        if 'mirror' in n: return 'Mirror'
        if 'mount' in n or 'charger' in n: return 'Mounts and Chargers'
        return 'Handlebars'
        
    if sheet_name == 'Foot Rests':
        return 'Foot Pegs and Mounts'
        
    if sheet_name == 'Guards & Mounts':
        if 'headlight' in n: return 'Head Light Grill'
        if 'master cylinder' in n: return 'Master Cylinder Guard'
        if 'fog' in n: return 'Fog Light Mount'
        if 'crash' in n: return 'Crash Guard'
        if 'radiator' in n: return 'Radiator Grill'
        return 'Crash Guard' # fallback
        
    if sheet_name == 'Oil Caps & Stands':
        if 'front' in n and 'cap' in n: return 'Front Oil Cap Silver'
        if 'rear' in n and 'cap' in n: return 'Rear Oil Cap'
        if 'side stand' in n: return 'Side Stand Base'
        if 'carrier' in n or 'rack' in n: return 'Luggage Carrier'
        if 'saddle' in n: return 'Saddle Stay'
        if 'cap' in n: return 'Front Oil Cap Silver'
        return 'Side Stand Base' # fallback
        
    if sheet_name == 'Body Fairing and Fenders':
        if 'tail' in n: return 'Tail'
        if 'fender' in n or 'mudguard' in n: return 'Fenders and Extenders'
        return 'Body Fairing'
        
    if sheet_name == 'Bike Essentials':
        if 'wind' in n or 'visor' in n: return 'Windshield'
        if 'winglet' in n: return 'Winglet'
        if 'seat' in n: return 'Seat'
        return 'Windshield'
        
    if sheet_name == 'Cleaning Accessories':
        return 'Cleaning Accessories'
        
    if sheet_name == 'Performance Accessories':
        if 'exhaust' in n: return 'Exhaust'
        if 'filter' in n: return 'Air Filter'
        if 'chain' in n or 'sprocket' in n: return 'Chain and Sprocket'
        if 'brake' in n or 'pad' in n or 'caliper' in n: return 'Braking Components'
        return 'Other Performance'
        
    if sheet_name == 'Lighting':
        if 'headlight' in n or 'bulb' in n: return 'Headlights'
        if 'indicator' in n: return 'Indicators'
        if 'fog' in n: return 'Fog Lights'
        if 'wire' in n or 'harness' in n or 'switch' in n or 'relay' in n: return 'Wiring and Electrical'
        if 'horn' in n: return 'Horns'
        if 'aux' in n: return 'Auxiliary Lights'
        return 'Other Lighting'
        
    if sheet_name == 'Riding Gears':
        if 'jacket' in n: return 'Riding Jackets'
        if 'glove' in n: return 'Gloves'
        if 'pant' in n: return 'Riding Pants'
        if 'boot' in n or 'shoe' in n: return 'Boots'
        if 'knee' in n or 'elbow' in n or 'bionic' in n or 'armor' in n or 'guard' in n: return 'Protective Armour'
        if 'suit' in n: return 'Riding Suits'
        return 'Other Riding Gear'
        
    if sheet_name == 'Luggage & Touring':
        if 'saddle' in n and 'stay' in n: return 'Saddle Stays'
        if 'saddle' in n: return 'Saddle Bags'
        if 'pannier' in n: return 'Panniers'
        if 'top box' in n or 'topcase' in n or 'top case' in n: return 'Top Boxes'
        if 'tank' in n: return 'Tank Bags'
        if 'rack' in n or 'carrier' in n: return 'Luggage Racks'
        if 'backrest' in n or 'bungee' in n: return 'Touring Accessories'
        return 'Other Luggage'
        
    if sheet_name == 'Helmets & Accessories':
        if 'full' in n: return 'Full-Face Helmet'
        if 'open' in n: return 'Open-Face Helmet'
        if 'modular' in n: return 'Modular Helmet'
        if 'visor' in n or 'shield' in n or 'lens' in n or 'pinlock' in n: return 'Helmet Visor'
        if 'helmet' in n: return 'Other Helmet'
        return 'Helmet Accessories'
        
    if sheet_name == 'General Accessories':
        if 'mobile' in n or 'phone' in n or 'gps' in n: return 'Mobile Holders'
        if 'charger' in n: return 'Chargers'
        if 'cover' in n: return 'Bike Covers'
        if 'seat' in n: return 'Seat Accessories'
        if 'lock' in n or 'alarm' in n: return 'Security Accessories'
        return 'Miscellaneous'

    return 'Miscellaneous'

main_sheets = wb.sheetnames[:14]
updated_count = 0

for s_name in main_sheets:
    sheet = wb[s_name]
    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row, column=1).value
        if not name: continue
        
        # Calculate strict dropdown subcategory
        new_subcat = determine_subcategory_strict(s_name, name)
        
        # Update Column 2
        sheet.cell(row=row, column=2, value=new_subcat)
        updated_count += 1

wb.save(file_path)
print(f"Strict dropdown subcategories fixed. Updated {updated_count} rows.")
