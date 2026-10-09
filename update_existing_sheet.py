import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

bikes = {
    "ROYAL ENFIELD": ["Scram 440 new", "Classic 650", "Bear 650", "Guerrilla 450", "Himalayan 450", "SUPER METEOR 650", "Hunter 350", "Himalayan", "Classic 350 Reborn", "Standard 350 Reborn", "Meteor 350", "Classic", "Standard", "Electra", "Interceptor", "Thunderbird", "Thunderbird X", "Continental GT", "Scram"],
    "YAMAHA": ["XSR 155", "RX 100", "Aerox 155", "MT 15", "R3", "R15 V1", "R15 V2", "R15 V3", "R15 V4", "R1"],
    "KTM": ["KTM 390 Enduro R", "Adventure 890", "ADVENTURE 390 2025 Model", "Adventure 390", "Duke 125", "Duke 200", "Duke 250 BS6", "Duke 250", "Duke 390", "RC 390"],
    "BAJAJ": ["NS400Z", "Pulsar 220", "Pulsar NS 200", "Pulsar RS 200", "Dominar 400", "Pulsar 150", "Pulsar 180"],
    "HONDA": ["XL750 Transalp", "NX500", "CB 200X", "Hness CB 350", "Hornet 160R", "CBR 250 R", "CB 300", "CB 350 RS", "CB 500X", "NX400", "CB650R", "CBR 650R", "CBR 1000RR"],
    "SUZUKI": ["Gixxer", "Gixxer SF", "V Strom 250", "Hayabusa", "GSX-S750", "V Strom 800DE", "V Strom DL650"],
    "HERO": ["Xpulse 210", "Mavrick 440", "Xpulse 200", "Impulse"],
    "TRIUMPH": ["Tiger 900", "Tracker 400", "Scrambler 400 x", "Speed 400", "Tiger Sport 660", "Tiger", "Trident 660"],
    "TVS": ["Apache RTX 300 new", "tvs ronin", "Apache 200", "Apache RTR 160", "Apache RTR 310"],
    "HARLEY DAVIDSON": ["Harley 440X"],
    "OLA": ["Ola"],
    "BMW": ["G 310GS", "G 310 R", "S 1000R", "R 1200GS", "R 1300GS"],
    "KAWASAKI": ["Versys 650", "Ninja 300", "Z900", "Ninja 400", "Ninja ZX-10R", "Ninja ZX6R"],
    "BENELLI": ["Benelli 600", "TRK 502"],
    "DUCATI": ["Panigale V4"],
    "APRILA": ["RS660"]
}

file_path = 'HRZ_Pitstop_Client_Inventory_Clean.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb.active

# Check if SKU is already first column
if ws['A1'].value != "SKU":
    ws.insert_cols(1)
    ws['A1'] = "SKU"

# Find which columns are Make and Model
# Make should be E, Model should be F if SKU is A
# Let's search dynamically to be safe
make_col_letter = None
model_col_letter = None

for cell in ws[1]:
    if "Make" in str(cell.value):
        make_col_letter = cell.column_letter
    elif "Model" in str(cell.value):
        model_col_letter = cell.column_letter

print(f"Make column: {make_col_letter}, Model column: {model_col_letter}")

# Add hidden sheet for data
if 'Data_Hidden' in wb.sheetnames:
    del wb['Data_Hidden']
hidden_sheet = wb.create_sheet('Data_Hidden')
hidden_sheet.sheet_state = 'hidden'

brands = list(bikes.keys())
for i, brand in enumerate(brands, 1):
    hidden_sheet.cell(row=i, column=1, value=brand)

wb.create_named_range('BrandList', hidden_sheet, f'$A$1:$A${len(brands)}')

col_idx = 2
for brand, models in bikes.items():
    safe_brand_name = brand.replace(" ", "_").replace("-", "_")
    for row_idx, model in enumerate(models, 1):
        hidden_sheet.cell(row=row_idx, column=col_idx, value=model)
    
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    wb.create_named_range(safe_brand_name, hidden_sheet, f'${col_letter}$1:${col_letter}${len(models)}')
    col_idx += 1

# Add Data Validation for Make
dv_make = DataValidation(type="list", formula1='"BrandList"', allow_blank=True)
ws.add_data_validation(dv_make)
dv_make.add(f'{make_col_letter}2:{make_col_letter}1048576')

# Add Data Validation for Model
# openpyxl has issues with dynamic INDIRECT references in data validation across many rows, 
# but it generally works if we use the first row's relative reference.
dv_model = DataValidation(type="list", formula1=f'INDIRECT(SUBSTITUTE({make_col_letter}2, " ", "_"))', allow_blank=True)
ws.add_data_validation(dv_model)
dv_model.add(f'{model_col_letter}2:{model_col_letter}1048576')

# Save over the same file
wb.save(file_path)
print("Updated HRZ_Pitstop_Client_Inventory_Clean.xlsx successfully!")
